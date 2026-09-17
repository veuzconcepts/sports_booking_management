"""Excel (openpyxl) and PDF (reportlab) exporters for reports.

Each function returns raw bytes so the view can wrap them in an HttpResponse
with the right content type.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

from apps.settings_app.currency import decimals_for, get_default_currency


def _dp() -> int:
    """Decimal places of the active currency (0/2/3) for export formatting."""
    return decimals_for(get_default_currency())


def _money(value, dp=None) -> str:
    """Format a money value to the currency's decimals with thousands separators."""
    dp = _dp() if dp is None else dp
    try:
        return f"{float(value):,.{dp}f}"
    except (TypeError, ValueError):
        return str(value)


def _xlsx_money_format(dp=None) -> str:
    dp = _dp() if dp is None else dp
    return "#,##0" if dp == 0 else "#,##0." + ("0" * dp)


def table_to_xlsx(title: str, subtitle: str, headers: list, rows: list) -> bytes:
    """Generic single-table xlsx. `rows` is a list of tuples (native types)."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    if subtitle:
        ws["A2"] = subtitle
    header_row = 4
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h).font = Font(bold=True)
    money_fmt = _xlsx_money_format()
    for i, row in enumerate(rows, start=header_row + 1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=val)
            if isinstance(val, float):   # money columns -> currency decimals
                cell.number_format = money_fmt
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + c)].width = 24
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table_to_pdf(title: str, subtitle: str, headers: list, rows: list) -> bytes:
    """Generic single-table PDF: first column left-aligned, rest right-aligned."""
    buf = BytesIO()
    pdf = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    left, right = 20 * mm, width - 20 * mm
    n = len(headers)

    def xpos(i):  # right edge for numeric columns
        return right - (n - 1 - i) * 40 * mm

    def fmt(v):
        return _money(v) if isinstance(v, float) else str(v)

    y = [height - 25 * mm]
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(left, y[0], title)
    y[0] -= 8 * mm
    if subtitle:
        pdf.setFont("Helvetica", 10)
        pdf.setFillColor(colors.grey)
        pdf.drawString(left, y[0], subtitle)
        pdf.setFillColor(colors.black)
    y[0] -= 12 * mm

    def header():
        pdf.setFillColor(colors.HexColor("#0f172a"))
        pdf.rect(left, y[0] - 2 * mm, right - left, 8 * mm, fill=1, stroke=0)
        pdf.setFillColor(colors.white)
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawString(left + 2 * mm, y[0], str(headers[0]))
        for i in range(1, n):
            pdf.drawRightString(xpos(i), y[0], str(headers[i]))
        pdf.setFillColor(colors.black)
        y[0] -= 9 * mm

    header()
    pdf.setFont("Helvetica", 9)
    for row in rows:
        if y[0] < 25 * mm:
            pdf.showPage()
            y[0] = height - 25 * mm
            pdf.setFont("Helvetica-Bold", 9)
            header()
            pdf.setFont("Helvetica", 9)
        pdf.drawString(left + 2 * mm, y[0], fmt(row[0]))
        for i in range(1, n):
            pdf.drawRightString(xpos(i), y[0], fmt(row[i]))
        y[0] -= 6 * mm

    pdf.showPage()
    pdf.save()
    return buf.getvalue()


def revenue_to_xlsx(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"

    ws["A1"] = "Revenue report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{report['date_from']} → {report['date_to']}"
    money_fmt = _xlsx_money_format()
    ws["A4"] = "Gross"
    ws["B4"] = float(report["gross_revenue"]); ws["B4"].number_format = money_fmt
    ws["A5"] = "Refunded"
    ws["B5"] = float(report["refunded"]); ws["B5"].number_format = money_fmt
    ws["A6"] = "Net"
    ws["B6"] = float(report["net_revenue"]); ws["B6"].number_format = money_fmt
    ws["A6"].font = ws["B6"].font = Font(bold=True)

    header_row = 8
    ws.cell(row=header_row, column=1, value="Date").font = Font(bold=True)
    ws.cell(row=header_row, column=2, value="Gross").font = Font(bold=True)
    ws.cell(row=header_row, column=3, value="Net").font = Font(bold=True)
    for i, row in enumerate(report["series"], start=header_row + 1):
        ws.cell(row=i, column=1, value=row["date"])
        g = ws.cell(row=i, column=2, value=float(row["gross"])); g.number_format = money_fmt
        nt = ws.cell(row=i, column=3, value=float(row["net"])); nt.number_format = money_fmt

    for col in ("A", "B", "C"):
        ws.column_dimensions[col].width = 16

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def revenue_to_pdf(report: dict) -> bytes:
    buf = BytesIO()
    pdf = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 25 * mm

    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(20 * mm, y, "Revenue Report")
    y -= 8 * mm
    pdf.setFont("Helvetica", 10)
    pdf.setFillColor(colors.grey)
    pdf.drawString(20 * mm, y, f"{report['date_from']} to {report['date_to']}")
    pdf.setFillColor(colors.black)

    y -= 14 * mm
    for label, key, bold in [("Gross revenue", "gross_revenue", False),
                             ("Refunded", "refunded", False),
                             ("Net revenue", "net_revenue", True)]:
        pdf.setFont("Helvetica-Bold" if bold else "Helvetica", 11)
        pdf.drawString(22 * mm, y, label)
        pdf.drawRightString(120 * mm, y, _money(report[key]))
        y -= 7 * mm

    y -= 6 * mm
    pdf.setFillColor(colors.HexColor("#0f172a"))
    pdf.rect(20 * mm, y - 2 * mm, 100 * mm, 8 * mm, fill=1, stroke=0)
    pdf.setFillColor(colors.white)
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawString(22 * mm, y, "Date")
    pdf.drawRightString(118 * mm, y, "Net")
    pdf.setFillColor(colors.black)
    y -= 9 * mm

    pdf.setFont("Helvetica", 9)
    for row in report["series"]:
        if y < 25 * mm:
            pdf.showPage()
            y = height - 25 * mm
            pdf.setFont("Helvetica", 9)
        pdf.drawString(22 * mm, y, row["date"])
        pdf.drawRightString(118 * mm, y, _money(row["net"]))
        y -= 6 * mm

    pdf.showPage()
    pdf.save()
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# AI Insights reports                                                          #
# --------------------------------------------------------------------------- #
# These render the dataset the assistant already produced. They never call the
# model and never recompute a figure: the file has to contain exactly what the
# user saw on screen.

_WIDGET_TITLE_FALLBACK = "Report"


def _sheet_name(title: str, used: set) -> str:
    """Excel sheet names are capped at 31 characters and must be unique."""
    base = (title or _WIDGET_TITLE_FALLBACK)[:31].strip() or _WIDGET_TITLE_FALLBACK
    name, n = base, 2
    while name.lower() in used:
        suffix = " %d" % n
        name = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def _as_number(value):
    """Money and counts arrive as strings; write them as numbers so Excel can
    total and chart them."""
    if isinstance(value, (int, float)):
        return value
    try:
        text = str(value)
        return float(text) if "." in text else int(text)
    except (TypeError, ValueError):
        return value


def ai_report_sheets(entry: dict) -> list:
    """Flatten the stored widgets into tabular sheets.

    One sheet per table or chart, plus the KPI blocks. A chart's underlying
    rows are what belongs in a spreadsheet; the picture is the screen's job.
    """
    sheets = []
    for widget in entry.get("widgets") or []:
        kind = widget.get("type")
        title = widget.get("title") or _WIDGET_TITLE_FALLBACK

        if kind == "kpi":
            sheets.append({
                "title": title,
                "headers": ["Metric", "Value"],
                "rows": [(item.get("label"), _as_number(item.get("value")))
                         for item in widget.get("items") or []],
            })
        elif kind == "table":
            columns = widget.get("columns") or []
            sheets.append({
                "title": title,
                "headers": [c.get("label") or c.get("key") for c in columns],
                "rows": [tuple(_as_number(row.get(c.get("key"))) for c in columns)
                         for row in widget.get("rows") or []],
            })
        elif kind in ("line", "bar", "donut"):
            x, y = widget.get("x"), widget.get("y")
            sheets.append({
                "title": title,
                "headers": [str(x).title(), str(y).title()],
                "rows": [(row.get(x), _as_number(row.get(y)))
                         for row in widget.get("rows") or []],
            })
    return sheets


def _report_subtitle(entry: dict) -> str:
    """Scope and period, so a file away from the screen still explains itself."""
    metas = [r.get("meta") or {} for r in entry.get("results") or []]
    period = ""
    for meta in metas:
        if meta.get("date_from") and meta.get("date_to"):
            period = "%s to %s" % (meta["date_from"], meta["date_to"])
            break
    scope = next((m.get("scope") for m in metas if m.get("scope")), "")
    return " | ".join([p for p in (period, scope) if p])


def ai_report_to_xlsx(entry: dict, sheets: list) -> bytes:
    """A workbook: cover sheet with the question and answer, then the data."""
    wb = Workbook()
    cover = wb.active
    cover.title = "Summary"
    cover["A1"] = "AI Insights report"
    cover["A1"].font = Font(bold=True, size=14)
    for row, (label, value) in enumerate((
        ("Question", entry.get("question", "")),
        ("Summary", entry.get("answer", "")),
        ("Scope", _report_subtitle(entry)),
        ("Generated", entry.get("stored_at", "")),
    ), start=3):
        cover.cell(row=row, column=1, value=label).font = Font(bold=True)
        cover.cell(row=row, column=2, value=value)
    cover.column_dimensions["A"].width = 18
    cover.column_dimensions["B"].width = 96
    cover["B4"].alignment = Alignment(wrap_text=True, vertical="top")

    used = {"summary"}
    money_fmt = _xlsx_money_format()
    for sheet in sheets:
        ws = wb.create_sheet(_sheet_name(sheet["title"], used))
        ws["A1"] = sheet["title"]
        ws["A1"].font = Font(bold=True, size=12)
        for column, header in enumerate(sheet["headers"], start=1):
            ws.cell(row=3, column=column, value=header).font = Font(bold=True)
        for index, row in enumerate(sheet["rows"], start=4):
            for column, value in enumerate(row, start=1):
                cell = ws.cell(row=index, column=column, value=value)
                if isinstance(value, float):
                    cell.number_format = money_fmt
        for column in range(1, len(sheet["headers"]) + 1):
            ws.column_dimensions[chr(64 + column)].width = 22
        # The header stays put while a long table scrolls.
        ws.freeze_panes = "A4"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ai_report_to_pdf(entry: dict, sheets: list) -> bytes:
    """A management document: title, scope, the summary, then each table."""
    buf = BytesIO()
    pdf = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    left, right = 20 * mm, width - 20 * mm
    y = height - 20 * mm

    def wrapped(text, font, size, max_width):
        """Break a paragraph to the page width; reportlab does not wrap."""
        words, line, lines = str(text or "").split(), "", []
        for word in words:
            candidate = (line + " " + word).strip()
            if pdf.stringWidth(candidate, font, size) <= max_width:
                line = candidate
            else:
                lines.append(line)
                line = word
        if line:
            lines.append(line)
        return lines

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(left, y, "AI Insights report")
    y -= 7 * mm

    subtitle = _report_subtitle(entry)
    if subtitle:
        pdf.setFont("Helvetica", 9)
        pdf.setFillColor(colors.grey)
        pdf.drawString(left, y, subtitle)
        pdf.setFillColor(colors.black)
        y -= 6 * mm

    pdf.setFont("Helvetica-Oblique", 9)
    pdf.setFillColor(colors.grey)
    pdf.drawString(left, y, "Generated " + str(entry.get("stored_at", "")))
    pdf.setFillColor(colors.black)
    y -= 9 * mm

    for label, value in (("Question", entry.get("question")),
                         ("Summary", entry.get("answer"))):
        if not value:
            continue
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawString(left, y, label)
        y -= 5 * mm
        pdf.setFont("Helvetica", 10)
        for line in wrapped(value, "Helvetica", 10, right - left):
            pdf.drawString(left, y, line)
            y -= 5 * mm
        y -= 3 * mm

    for sheet in sheets:
        if y < 45 * mm:
            pdf.showPage()
            y = height - 20 * mm

        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawString(left, y, sheet["title"])
        y -= 6 * mm

        headers, rows = sheet["headers"], sheet["rows"]
        col_width = (right - left) / max(1, len(headers))

        pdf.setFont("Helvetica-Bold", 9)
        for index, header in enumerate(headers):
            x = left + index * col_width
            if index == 0:
                pdf.drawString(x, y, str(header))
            else:
                pdf.drawRightString(x + col_width - 2 * mm, y, str(header))
        y -= 2 * mm
        pdf.setStrokeColor(colors.lightgrey)
        pdf.line(left, y, right, y)
        y -= 4 * mm

        pdf.setFont("Helvetica", 9)
        for row in rows:
            if y < 20 * mm:
                pdf.showPage()
                y = height - 20 * mm
                pdf.setFont("Helvetica", 9)
            for index, value in enumerate(row):
                x = left + index * col_width
                text = _money(value) if isinstance(value, float) else str(value)
                if index == 0:
                    pdf.drawString(x, y, text[:48])
                else:
                    pdf.drawRightString(x + col_width - 2 * mm, y, text)
            y -= 5 * mm
        y -= 6 * mm

    pdf.showPage()
    pdf.save()
    return buf.getvalue()
